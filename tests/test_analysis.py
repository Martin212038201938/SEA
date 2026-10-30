import datetime as dt

from cockpit.analysis import anomalies, experiments, kpis
from cockpit.models import Experiment
from cockpit.util import yesterday


def test_overview_sums(session):
    end = yesterday()
    ov = kpis.overview(session, end - dt.timedelta(days=27), end, "google")
    assert ov["cost"] > 0 and ov["clicks"] > 0
    assert ov["avg_cpc"] == ov["cost"] / ov["clicks"]
    assert ov["days"] == 28


def test_compare_periods_has_deltas(session):
    cmp = kpis.compare_periods(session, 14, "all")
    assert set(cmp["deltas"]) >= {"cost", "clicks", "cpl"}
    assert cmp["previous"]["cost"] >= 0


def test_search_terms_filter(session):
    end = yesterday()
    waste = kpis.search_terms(session, end - dt.timedelta(days=27), end, min_cost=1, only_without_conversions=True)
    assert all(t["conversions"] == 0 for t in waste)
    assert all(t["cost"] >= 1 for t in waste)


def test_anomalies_do_not_duplicate(session):
    day = yesterday()
    cands = anomalies.evaluate(session, day, "google")
    created1 = anomalies.persist(session, day, "google", cands)
    created2 = anomalies.persist(session, day, "google", cands)
    assert len(created1) == len(cands) or len(created1) <= len(cands)
    assert created2 == []


def test_experiment_compare_windows(session):
    e = session.query(Experiment).first()
    cmp = experiments.compare(session, e)
    w = cmp["windows"]
    assert w["base_end"] < w["test_start"]
    base_len = (dt.date.fromisoformat(w["base_end"]) - dt.date.fromisoformat(w["base_start"])).days
    test_len = (dt.date.fromisoformat(w["test_end"]) - dt.date.fromisoformat(w["test_start"])).days
    assert base_len == test_len
    assert "verdict" in cmp and cmp["verdict"]


def test_wilson_interval():
    lo, hi = experiments.wilson_interval(3, 100)
    assert 0 < lo < 0.03 < hi < 0.1
    assert experiments.wilson_interval(0, 0) is None


def test_zero_impressions_plausibility_and_counting(session):
    import datetime as dt
    from cockpit.analysis import anomalies
    from cockpit.models import CampaignDaily, ConversionAction

    day = dt.date(2026, 10, 27)  # Dienstag
    for i in range(1, 10):
        d = day - dt.timedelta(days=i)
        session.add(CampaignDaily(date=d, platform="google", campaign_id="P1", campaign_name="Plausi", status="ENABLED", impressions=100, clicks=5, cost=40.0))
    # Import unvollständig: Kosten und Klicks da, Impressionen 0 -> kein Alarm
    row = CampaignDaily(date=day, platform="google", campaign_id="P1", campaign_name="Plausi", status="ENABLED", impressions=0, clicks=4, cost=148.0)
    session.add(row)
    session.add(ConversionAction(platform="google", external_id="CA9", name="Lead Formular senden", category="SUBMIT_LEAD_FORM", status="ENABLED", primary_for_goal=True, counting_type="MANY_PER_CLICK"))
    session.flush()
    rules = {(a["rule"], a["key"]) for a in anomalies.evaluate(session, day, "google")}
    assert ("zero_impressions", "P1") not in rules
    assert ("counting_type", "CA9") in rules
    # echter Ausfall: alles 0 -> Alarm
    row.clicks, row.cost = 0, 0.0
    session.flush()
    assert ("zero_impressions", "P1") in {(a["rule"], a["key"]) for a in anomalies.evaluate(session, day, "google")}
    assert anomalies.evaluation_day(session, dt.date(2026, 10, 28)) == dt.date(2026, 10, 26)
    session.rollback()


def test_sea_excel_parse_and_kpi(session, tmp_path):
    import datetime as dt
    import io
    from openpyxl import Workbook
    from cockpit.connectors import sea_excel
    from cockpit.analysis import kpis

    wb = Workbook(); ws = wb.active
    ws["B2"], ws["C2"] = dt.datetime(2026, 9, 7), "13.10.2026"
    ws["B6"], ws["B8"], ws["B10"], ws["B12"], ws["B18"] = "Vorname: Kim", "Nachname: Tretter", "E-Mail: ktretter@rajapack.de", "Unternehmen: Rajapack GmbH", "https://trello.com/c/x"
    ws["C10"], ws["C12"] = "E-Mail: a@b.de", "Unternehmen: B GmbH"
    ws["D10"] = ""  # leere Spalte wird übersprungen
    ws["A20"], ws["A21"] = "Status", "Erlös"
    ws["B20"], ws["C20"] = "7.600", None
    ws["B21"] = 1520
    buf = io.BytesIO(); wb.save(buf)
    rows = sea_excel.parse(buf.getvalue())
    assert [r["column"] for r in rows] == ["B", "C"]
    assert rows[0]["net_revenue"] == 7600.0 and rows[0]["margin"] == 1520.0 and rows[0]["company"] == "Rajapack GmbH" and rows[0]["date"] == dt.date(2026, 9, 7)
    assert rows[1]["date"] == dt.date(2026, 10, 13) and rows[1]["net_revenue"] is None
    sea_excel.store(session, rows)
    k = kpis.sea_revenue(session, dt.date(2026, 9, 1), dt.date(2026, 10, 31))
    assert k["ads_leads"] == 2 and k["net_revenue"] == 7600.0 and k["leads_with_revenue"] == 1
    assert sea_excel._num("7.600,50") == 7600.5 and sea_excel._num("1520") == 1520.0
    session.rollback()


def test_sea_excel_steps_a_b_idempotent(session):
    import datetime as dt
    import io
    from openpyxl import Workbook
    from cockpit.connectors import sea_excel
    from cockpit.models import Lead, RevenueLog

    wb = Workbook(); ws = wb.active
    ws["B2"], ws["B10"], ws["B12"], ws["B18"] = dt.datetime(2026, 8, 20), "E-Mail: bach@aps-payment.de", "Unternehmen: APS", "https://trello.com/c/yaC4EpvK"
    ws["C2"], ws["C10"], ws["C12"] = dt.datetime(2026, 9, 7), "E-Mail: kim@rajapack.de", "Unternehmen: Rajapack"
    ws["D2"], ws["D10"] = dt.datetime(2026, 9, 9), "E-Mail: niemand@x.de"
    ws["A20"], ws["A21"], ws["A23"], ws["A24"], ws["A25"], ws["A26"] = "Status", "Erlös", "Auftragswert (gewonnen)", "Beleg A", "Umsatz final (erledigt)", "Beleg B"
    ws["B20"], ws["B21"] = 4500, "=B20*0.2"  # Formel ohne gespeicherten Wert, wie nach dem Speichern per Skript
    ws["B23"], ws["B24"], ws["B25"], ws["B26"] = 5000, "AG0042 vom 01.09.2026", 4500, "RE0100, RE0101"
    ws["C23"], ws["C24"] = 2800, "AG0050"
    ws["D23"] = 1000
    buf = io.BytesIO(); wb.save(buf)
    rows = sea_excel.parse(buf.getvalue())
    b = rows[0]
    assert (b["order_value"], b["order_doc"], b["final_revenue"], b["final_doc"]) == (5000.0, "AG0042 vom 01.09.2026", 4500.0, "RE0100, RE0101")
    assert b["margin"] == 900.0  # aus der Formel nachgerechnet, nicht der veraltete Wert
    sea_excel.store(session, rows)

    aps = Lead(source="trello", external_id="c1", created_on=dt.date(2026, 8, 21), status="kontaktiert", url="https://trello.com/c/yaC4EpvK")
    raja = Lead(source="trello", external_id="c2", created_on=dt.date(2026, 9, 9), status="abgeschlossen", email="kim@rajapack.de")
    session.add_all([aps, raja]); session.flush()

    dry = sea_excel.apply_to_leads(session, dry_run=True)
    assert dry["gesetzt"] == 3 and aps.value == 0 and not aps.value_ref and session.query(RevenueLog).count() == 0

    st = sea_excel.apply_to_leads(session)
    assert st == {"gesetzt": 3, "geändert": 0, "entfernt": 0, "kein-lead": 1}
    assert (aps.value, aps.revenue, raja.value, raja.revenue) == (5000.0, 4500.0, 2800.0, 0.0)
    assert aps.status == "gewonnen" and raja.status == "gewonnen"  # Status folgt Lexoffice, nicht Trello
    assert aps.value_ref.startswith("excel:B A AG0042") and aps.revenue_ref.startswith("excel:B B RE0100")

    # zweiter Lauf ohne Änderung: nichts zu tun
    assert sea_excel.apply_to_leads(session) == {"gesetzt": 0, "geändert": 0, "entfernt": 0, "kein-lead": 1}
    # Korrektur in der Excel wird nachgezogen und protokolliert
    session.query(sea_excel.SeaRevenue).filter_by(column="B").one().final_revenue = 4300
    assert sea_excel.apply_to_leads(session)["geändert"] == 1 and aps.revenue == 4300.0
    assert session.query(RevenueLog).filter_by(action="geändert").count() == 1
    # Wert in der Excel gelöscht (z. B. Storno) → Lead-Wert entfernt
    session.query(sea_excel.SeaRevenue).filter_by(column="C").one().order_value = None
    assert sea_excel.apply_to_leads(session)["entfernt"] == 1 and raja.value == 0.0 and raja.value_ref == ""
    session.rollback()


def test_trello_sync_keeps_locked_values(session, monkeypatch):
    import datetime as dt
    from cockpit.connectors import trello
    from cockpit.models import Lead

    session.add(Lead(source="trello", external_id="66d4a1b2c3d4e5f6a7b8c9d0", created_on=dt.date(2026, 9, 1), status="gewonnen", value=5000, value_ref="excel:B A AG0042", revenue=0, revenue_ref=""))
    session.add(Lead(source="trello", external_id="66d4a1b2c3d4e5f6a7b8c9d1", created_on=dt.date(2026, 9, 1), status="neu"))
    session.flush()
    lists = [{"id": "L1", "name": "Erledigt"}]
    cards = [{"id": "66d4a1b2c3d4e5f6a7b8c9d0", "name": "APS", "desc": "Auftragswert: 100\nErlös: 50", "idList": "L1", "shortUrl": "https://trello.com/c/x", "labels": [], "closed": False},
             {"id": "66d4a1b2c3d4e5f6a7b8c9d1", "name": "Rajapack", "desc": "", "idList": "L1", "shortUrl": "https://trello.com/c/y", "labels": [], "closed": False}]
    monkeypatch.setattr(trello, "_get", lambda path, **kw: lists if path.endswith("/lists") else cards)
    monkeypatch.setattr(trello.settings, "trello_list_status_map", {"Erledigt": "abgeschlossen"})
    trello.pull(session)
    lead = session.query(Lead).filter_by(external_id="66d4a1b2c3d4e5f6a7b8c9d0").one()
    assert lead.value == 5000 and lead.revenue == 50 and lead.status == "gewonnen"  # Lexoffice-Werte und Status bleiben
    other = session.query(Lead).filter_by(external_id="66d4a1b2c3d4e5f6a7b8c9d1").one()
    assert other.status == "abgeschlossen"  # „Erledigt“ heißt nicht gewonnen
    session.rollback()
