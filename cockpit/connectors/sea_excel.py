"""Excel „Conversions SEA“ (SharePoint) lesen: je Spalte ein Lead aus der Anzeigenwerbung mit Nettoumsatz.

Aufbau (spaltenweise ab B): Zeile 2 Datum, 6 Vorname, 8 Nachname, 10 E-Mail, 12 Unternehmen, 18 Trello-Link,
Zeile mit Beschriftung „Status“ in Spalte A = Nettoumsatz, Zeile „Erlös“ = Marge (Formel). Dazu die Zeilen, die die
Sales-Assistant-Aufgabe pflegt: „Auftragswert (gewonnen)“ und „Beleg A“ (Schritt A, Karte gewonnen, angenommenes
Angebot), „Umsatz final (erledigt)“ und „Beleg B“ (Schritt B, Karte erledigt, bezahlte Rechnungen). Alle diese Zeilen
werden über die Beschriftung in Spalte A gesucht, damit eingefügte Zeilen nichts kaputt machen.
Die Marge rechnet das Cockpit aus der Formel selbst nach: Speichert ein Skript die Datei, fehlen sonst neu berechnete
Formelwerte und es stünde der alte Wert drin.
Übernahme in die Leads (apply_to_leads): Zuordnung über den Trello-Link, sonst die E-Mail. Jeder Wert wird genau einmal
gesetzt und dann gesperrt (Lead.value_ref / revenue_ref); spätere Abweichungen werden nur protokolliert.
Zugriff: Microsoft Graph, App-only (Client Credentials), Berechtigung Sites.Selected mit Leserecht auf die eine Site.
"""
from __future__ import annotations

import datetime as dt
import io
import logging
import re
from typing import Any
from urllib.parse import quote

import httpx
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..config import settings
from ..models import Lead, RevenueLog, SeaRevenue, utcnow

log = logging.getLogger(__name__)
GRAPH = "https://graph.microsoft.com/v1.0"
FIELD_ROWS = {"date": 2, "first": 6, "last": 8, "email": 10, "company": 12, "trello": 18}


def _token() -> str:
    r = httpx.post(
        f"https://login.microsoftonline.com/{settings.graph_tenant_id}/oauth2/v2.0/token",
        data={"grant_type": "client_credentials", "client_id": settings.graph_client_id, "client_secret": settings.graph_client_secret, "scope": "https://graph.microsoft.com/.default"},
        timeout=30,
    )
    r.raise_for_status()
    return r.json()["access_token"]


def _get(client: httpx.Client, url: str) -> httpx.Response:
    r = client.get(url, timeout=60, follow_redirects=True)
    r.raise_for_status()
    return r


def download() -> bytes:
    with httpx.Client(headers={"Authorization": f"Bearer {_token()}"}) as c:
        site = _get(c, f"{GRAPH}/sites/{settings.sea_excel_site}").json()
        item = f"{GRAPH}/sites/{site['id']}/drive/root:/{quote(settings.sea_excel_path)}:/content"
        return _get(c, item).content


def _strip_label(v: Any) -> str:
    s = "" if v is None else str(v).strip()
    return re.sub(r"^(vorname|nachname|e-?mail|unternehmen|telefon)\s*:\s*", "", s, flags=re.I).strip()


def _num(v: Any) -> float | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^\d,.\-]", "", str(v))
    if not s:
        return None
    if "," in s:  # deutsches Format 7.600,50
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):  # 7.600
        s = s.replace(".", "")
    try:
        return float(s)
    except ValueError:
        return None


def _date(v: Any) -> dt.date | None:
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    s = str(v or "").strip()
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%m/%d/%Y"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _label_rows(ws) -> dict[str, int]:
    out: dict[str, int] = {}
    for r in range(1, ws.max_row + 1):
        label = re.sub(r"\s+", " ", str(ws.cell(row=r, column=1).value or "").strip().lower())
        if label and label not in out:
            out[label] = r
    return out


def _find(labels: dict[str, int], *prefixes: str) -> int | None:
    for p in prefixes:
        for label, row in labels.items():
            if label.startswith(p):
                return row
    return None


def _margin(formula: Any, cached: Any, net: float | None) -> float | None:
    """Marge aus der Formel (=B20*0.2) nachrechnen; nur ohne erkennbare Formel den gespeicherten Wert nehmen."""
    m = re.fullmatch(r"=\s*\$?[A-Z]+\$?\d+\s*\*\s*([\d.,]+)", str(formula or "").strip())
    if m:
        factor = _num(m.group(1))
        return round(net * factor, 2) if (net is not None and factor is not None) else None
    return _num(cached)


def parse(content: bytes) -> list[dict[str, Any]]:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    ws = load_workbook(io.BytesIO(content), data_only=True).worksheets[0]
    wf = load_workbook(io.BytesIO(content), data_only=False).worksheets[0]
    labels = _label_rows(ws)
    status_row = labels.get("status")
    margin_row = labels.get("erlös") or labels.get("erloes")
    order_row = _find(labels, "auftragswert")
    order_doc_row = _find(labels, "beleg a")
    final_row = _find(labels, "umsatz final")
    final_doc_row = _find(labels, "beleg b")

    def cell(row: int | None, col: int, sheet=ws) -> Any:
        return sheet.cell(row=row, column=col).value if row else None

    out = []
    for col in range(2, ws.max_column + 1):
        email = _strip_label(ws.cell(row=FIELD_ROWS["email"], column=col).value).lower()
        if "@" not in email:
            continue
        name = " ".join(x for x in (_strip_label(ws.cell(row=FIELD_ROWS["first"], column=col).value), _strip_label(ws.cell(row=FIELD_ROWS["last"], column=col).value)) if x)
        net = _num(cell(status_row, col))
        out.append(
            {
                "column": get_column_letter(col),
                "date": _date(ws.cell(row=FIELD_ROWS["date"], column=col).value),
                "name": name,
                "email": email,
                "company": _strip_label(ws.cell(row=FIELD_ROWS["company"], column=col).value),
                "trello_url": _strip_label(ws.cell(row=FIELD_ROWS["trello"], column=col).value),
                "net_revenue": net,
                "margin": _margin(cell(margin_row, col, wf), cell(margin_row, col), net) if margin_row else None,
                "order_value": _num(cell(order_row, col)),
                "order_doc": str(cell(order_doc_row, col) or "").strip()[:255],
                "final_revenue": _num(cell(final_row, col)),
                "final_doc": str(cell(final_doc_row, col) or "").strip()[:255],
            }
        )
    return out


def store(session: Session, rows: list[dict[str, Any]]) -> int:
    existing = {r.column: r for r in session.execute(select(SeaRevenue)).scalars()}
    seen = set()
    for row in rows:
        obj = existing.get(row["column"]) or SeaRevenue(column=row["column"])
        for k, v in row.items():
            setattr(obj, k, v)
        obj.synced_at = utcnow()
        session.add(obj)
        seen.add(row["column"])
    for col, obj in existing.items():
        if col not in seen:
            session.delete(obj)
    session.flush()
    return len(rows)


def _short(url: str) -> str:
    m = re.search(r"trello\.com/c/([A-Za-z0-9]+)", url or "")
    return m.group(1) if m else ""


def _lead_for(session: Session, row: SeaRevenue) -> Lead | None:
    short = _short(row.trello_url)
    if short:
        for lead in session.execute(select(Lead).where(Lead.url.like(f"%/c/{short}%"))).scalars():
            return lead
    if row.email:
        return session.execute(select(Lead).where(Lead.email == row.email).order_by(Lead.created_on)).scalars().first()
    return None


def apply_to_leads(session: Session, dry_run: bool = False, log_fn=None) -> dict[str, int]:
    """Schritt A (Auftragswert) und B (finaler Umsatz) aus der Excel in die Leads übernehmen, je einmal.
    dry_run: nur berichten, nichts schreiben."""
    log_fn = log_fn or (lambda s: log.info(s))
    stats = {"gesetzt": 0, "abweichung": 0, "kein-lead": 0}
    prefix = "TROCKENLAUF " if dry_run else ""
    for row in session.execute(select(SeaRevenue).order_by(SeaRevenue.column)).scalars():
        steps = [("A", row.order_value, row.order_doc, "value", "value_ref"), ("B", row.final_revenue, row.final_doc, "revenue", "revenue_ref")]
        if not any(v for _, v, *_ in steps):
            continue
        lead = _lead_for(session, row)
        for step, amount, doc, field, ref_field in steps:
            if not amount:
                continue
            ref = f"excel:{row.column} {step} {doc}".strip()[:255]
            if lead is None:
                action, note = "kein-lead", f"Kein Lead zu Spalte {row.column} ({row.company}, {row.trello_url or row.email})"
            elif not getattr(lead, ref_field):
                action, note = "gesetzt", f"{field} {getattr(lead, field) or 0:.2f} → {amount:.2f}"
                if not dry_run:
                    setattr(lead, field, amount)
                    setattr(lead, ref_field, ref)
            elif abs((getattr(lead, field) or 0) - amount) > 0.005:
                action, note = "abweichung", f"{field} bleibt {getattr(lead, field):.2f} ({getattr(lead, ref_field)}), Excel sagt {amount:.2f}"
            else:
                continue  # schon übernommen, nichts zu tun
            stats[action] += 1
            log_fn(f"{prefix}Umsatz {step} Lead {lead.id if lead else '-'} Spalte {row.column} {row.company}: {action}, {amount:.2f} €, Beleg {doc or '-'} – {note}")
            if not dry_run:
                session.add(RevenueLog(lead_id=lead.id if lead else None, column=row.column, step=step, amount=amount, doc=doc, action=action, note=note))
    if not dry_run:
        session.flush()
    return stats


def pull(session: Session, log_fn=None, dry_run: bool = False) -> int:
    log_fn = log_fn or (lambda s: log.info(s))
    rows = parse(download())
    n = store(session, rows)
    log_fn(f"Excel Conversions SEA: {n} Leads, Umsatz gesamt {sum(r['net_revenue'] or 0 for r in rows):.0f} €")
    stats = apply_to_leads(session, dry_run=dry_run, log_fn=log_fn)
    log_fn(f"Excel → Leads: {stats['gesetzt']} übernommen, {stats['abweichung']} Abweichungen, {stats['kein-lead']} ohne Lead")
    return n


def test_connection() -> str:
    rows = parse(download())
    return f"Gelesen: {len(rows)} Leads, davon {sum(1 for r in rows if r['net_revenue'])} mit Umsatz"
